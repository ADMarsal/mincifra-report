"""
Шаг 1. Чтение общей таблицы стажёров из Excel.

Результат - список словарей, по одному словарю на студента:
    {"name": "Студент 01", "university": "...", "qualification": "Бакалавр",
     "speciality": "Менеджмент", "hours": 640}
"""

from openpyxl import load_workbook

from config import COLUMNS, HEADER_MARKER


def find_header_row(sheet) -> int:
    """Вернуть номер строки, где стоит заголовок "Сотрудник"."""
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == HEADER_MARKER:
                return cell.row
    raise ValueError(f"В таблице не найдена строка с заголовком «{HEADER_MARKER}»")


def find_columns(sheet, header_row: int) -> dict:
    """Вернуть словарь: поле программы -> номер столбца в Excel."""
    columns = {}
    for cell in sheet[header_row]:
        title = str(cell.value).lower() if cell.value is not None else ""
        for field, word in COLUMNS.items():
            if word in title and field not in columns:
                columns[field] = cell.column

    missing = [field for field in COLUMNS if field not in columns]
    if missing:
        raise ValueError(f"В таблице нет нужных столбцов: {', '.join(missing)}")
    return columns


def read_students(path: str) -> list[dict]:
    """Прочитать файл Excel и вернуть список студентов (словарей)."""
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.worksheets[0]
    header_row = find_header_row(sheet)
    columns = find_columns(sheet, header_row)

    students = []
    for row in range(header_row + 1, sheet.max_row + 1):
        name = sheet.cell(row, columns["name"]).value
        if name is None:
            continue
        student = {}
        for field, column in columns.items():
            student[field] = sheet.cell(row, column).value
        students.append(student)
    return students
